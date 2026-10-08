#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Construye la PLANTILLA EN BLANCO de cotizaciones ITSA en formato Word (.docx).

Replica la maqueta de COT-2026-005_Impoventura_v1.pdf y deja todos los datos
como marcadores {{CAMPO}} para que la Gema (o cualquier script) los reemplace.

Uso:  python3 crear_plantilla_docx.py  [ruta_salida.docx]
"""
import sys, os
from docx import Document
from docx.shared import Pt, Mm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG = os.path.join(RAIZ, 'assets', 'img')

# ---- Tokens de marca (Brandbook ITSA) -------------------------------------
AZUL        = RGBColor(0x1C, 0x58, 0xD7)
AZUL_HEX    = '1C58D7'
GRIS_NUM    = RGBColor(0xC2, 0xC2, 0xC2)
GRIS_LABEL  = RGBColor(0x85, 0x85, 0x85)
GRIS_TEXTO  = RGBColor(0x5C, 0x5D, 0x5D)
NEGRO       = RGBColor(0x0C, 0x0C, 0x0C)
GRIS_FONDO  = 'F7F7F7'
GRIS_LINEA  = 'D6D6D6'
FUENTE      = 'Schibsted Grotesk'   # fallback automático a Arial si no está instalada


# ---- Utilidades XML --------------------------------------------------------
def _borde(tc_pr, lado, color, sz=6):
    bordes = tc_pr.find(qn('w:tcBorders'))
    if bordes is None:
        bordes = OxmlElement('w:tcBorders'); tc_pr.append(bordes)
    el = OxmlElement('w:' + lado)
    el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), '0');    el.set(qn('w:color'), color)
    bordes.append(el)


def bordes_celda(celda, color=AZUL_HEX, sz=6):
    tc_pr = celda._tc.get_or_add_tcPr()
    for lado in ('top', 'left', 'bottom', 'right'):
        _borde(tc_pr, lado, color, sz)


def fondo_celda(celda, hexa):
    tc_pr = celda._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear'); shd.set(qn('w:color'), 'auto'); shd.set(qn('w:fill'), hexa)
    tc_pr.append(shd)


def no_partir(fila):
    """Impide que Word corte la fila entre dos páginas (pie de firma)."""
    tr_pr = fila._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement('w:cantSplit'))


def limpiar(celda):
    """Deja una sola párrafo vacío: al combinar celdas Word acumula los de origen."""
    for extra in list(celda.paragraphs)[1:]:
        extra._element.getparent().remove(extra._element)
    celda.paragraphs[0].text = ''
    return celda.paragraphs[0]


def alto_minimo(fila, mm):
    tr_pr = fila._tr.get_or_add_trPr()
    h = OxmlElement('w:trHeight')
    h.set(qn('w:val'), str(int(mm * 56.7)))   # mm -> twips
    h.set(qn('w:hRule'), 'atLeast')
    tr_pr.append(h)


def linea_inferior(parrafo, color=GRIS_LINEA):
    p_pr = parrafo._p.get_or_add_pPr()
    bordes = OxmlElement('w:pBdr'); p_pr.append(bordes)
    b = OxmlElement('w:bottom')
    b.set(qn('w:val'), 'single'); b.set(qn('w:sz'), '4')
    b.set(qn('w:space'), '4');    b.set(qn('w:color'), color)
    bordes.append(b)


def espaciado_letras(run, twentieths):
    r_pr = run._element.get_or_add_rPr()
    sp = OxmlElement('w:spacing'); sp.set(qn('w:val'), str(twentieths))
    r_pr.append(sp)


def txt(parrafo, texto, *, size=10, bold=False, color=NEGRO, caps=False, tracking=None):
    run = parrafo.add_run(texto.upper() if caps else texto)
    run.font.name = FUENTE
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color
    r_pr = run._element.get_or_add_rPr()
    rf = r_pr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); r_pr.insert(0, rf)
    for a in ('w:ascii', 'w:hAnsi', 'w:cs'):
        rf.set(qn(a), FUENTE)
    if tracking:
        espaciado_letras(run, tracking)
    return run


def parrafo(doc_o_celda, texto='', **kw):
    espacio_antes = kw.pop('antes', None)
    espacio_despues = kw.pop('despues', 0)
    alineacion = kw.pop('align', None)
    p = doc_o_celda.add_paragraph()
    p.paragraph_format.space_before = Pt(espacio_antes if espacio_antes is not None else 0)
    p.paragraph_format.space_after = Pt(espacio_despues)
    if alineacion is not None:
        p.alignment = alineacion
    if texto:
        txt(p, texto, **kw)
    return p


def titulo_seccion(doc, numero, texto):
    p = parrafo(doc, antes=16, despues=6)
    txt(p, numero + ' ', size=13, bold=True, color=GRIS_NUM)
    txt(p, texto, size=13, bold=True, color=AZUL)
    return p


# ---- Construcción del documento -------------------------------------------
def construir(salida):
    doc = Document()

    est = doc.styles['Normal']
    est.font.name = FUENTE
    est.font.size = Pt(10)
    est.element.rPr.rFonts.set(qn('w:eastAsia'), FUENTE)

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Mm(210), Mm(297)
    sec.top_margin, sec.bottom_margin = Mm(18), Mm(16)
    sec.left_margin, sec.right_margin = Mm(20), Mm(20)
    sec.header_distance = Mm(10)

    # ---------- Encabezado (se repite en todas las páginas) ----------
    enc = sec.header
    enc.paragraphs[0].text = ''
    t = enc.add_table(rows=1, cols=2, width=Mm(170))
    t.autofit = False
    t.columns[0].width, t.columns[1].width = Mm(95), Mm(75)
    c_logo, c_etq = t.rows[0].cells
    c_logo.width, c_etq.width = Mm(95), Mm(75)
    p = c_logo.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(os.path.join(IMG, 'logo_itsanet_ecuador.png'), width=Mm(45))
    p = c_etq.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after = Pt(0)
    txt(p, '{{ETIQUETA_DOCUMENTO}}', size=7.5, bold=True, color=AZUL, caps=True, tracking=60)
    linea_inferior(enc.add_paragraph())

    # ---------- Apertura ----------
    parrafo(doc, '{{TITULO_DOCUMENTO}}', size=24, bold=True, color=AZUL, antes=6, despues=10)
    parrafo(doc, '{{CIUDAD_EMISION}}, {{FECHA_EMISION_LARGA}}',
            size=10.5, color=GRIS_TEXTO, align=WD_ALIGN_PARAGRAPH.RIGHT, despues=18)
    parrafo(doc, 'Atención a:', despues=10)
    parrafo(doc, 'Equipo de {{CLIENTE_MAYUSCULAS}}', bold=True, despues=10)
    parrafo(doc, 'Presente.-', despues=4)

    # ---------- 01 Datos ----------
    titulo_seccion(doc, '01', 'Datos de la cotización')
    campos = [('CLIENTE', '{{CLIENTE}}'),            ('ATENCIÓN A', '{{ATENCION_A}}'),
              ('N.º DE COTIZACIÓN', '{{CODIGO_DOCUMENTO}}'),
              ('FECHA DE EMISIÓN', '{{FECHA_EMISION_LARGA}}'),
              ('VALIDEZ DE LA OFERTA', '{{VALIDEZ}}'),
              ('GERENTE COMERCIAL', '{{GERENTE_COMERCIAL}}'),
              ('SERVICIO COTIZADO', '{{SERVICIO_COTIZADO}}'),
              ('CIUDAD(ES) DE OPERACIÓN', '{{CIUDADES_OPERACION}}')]
    t = doc.add_table(rows=4, cols=2); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, fila in enumerate(t.rows):
        alto_minimo(fila, 16)
        for j, celda in enumerate(fila.cells):
            celda.width = Mm(85)
            bordes_celda(celda)
            etiqueta, valor = campos[i * 2 + j]
            celda.paragraphs[0].text = ''
            p = celda.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
            txt(p, etiqueta, size=6.5, color=GRIS_LABEL, caps=True, tracking=24)
            parrafo(celda, valor, size=10, color=NEGRO)

    # ---------- 02 Detalle económico ----------
    titulo_seccion(doc, '02', 'Detalle económico')
    anchos = [Mm(64), Mm(38), Mm(34), Mm(34)]
    t = doc.add_table(rows=1, cols=4); t.alignment = WD_TABLE_ALIGNMENT.CENTER
    cab = ['Servicio', 'Cant.', 'Tarifa ({{MONEDA}})', 'Costo total estimado']
    for j, celda in enumerate(t.rows[0].cells):
        celda.width = anchos[j]; bordes_celda(celda); fondo_celda(celda, GRIS_FONDO)
        celda.paragraphs[0].text = ''
        p = celda.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        txt(p, cab[j], size=9.5, bold=True, color=AZUL)

    # Fila de ítem: es el patrón que se duplica tantas veces como servicios haya
    fila = t.add_row(); alto_minimo(fila, 16)
    valores = ['{{ITEM_SERVICIO}}', '{{ITEM_CANTIDAD}}', '{{ITEM_TARIFA}}', '{{ITEM_TOTAL}}']
    for j, celda in enumerate(fila.cells):
        celda.width = anchos[j]; bordes_celda(celda)
        celda.paragraphs[0].text = ''
        p = celda.paragraphs[0]; p.paragraph_format.space_after = Pt(0)
        if j == 3:
            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        txt(p, valores[j], size=9.5)

    # Fila de totales (opcional: se elimina si la cotización no los lleva)
    fila = t.add_row()
    for j, celda in enumerate(fila.cells):
        celda.width = anchos[j]; bordes_celda(celda); fondo_celda(celda, GRIS_FONDO)
        celda.paragraphs[0].text = ''
    izq = fila.cells[0].merge(fila.cells[2])
    p = limpiar(izq); p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after = Pt(0)
    txt(p, '{{TOTAL_ETIQUETA}}', size=9.5, bold=True, color=AZUL)
    p = fila.cells[-1].paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT; p.paragraph_format.space_after = Pt(0)
    txt(p, '{{TOTAL_VALOR}}', size=9.5, bold=True)

    # Fila "Otros:" + nota fiscal
    fila = t.add_row()
    for j, celda in enumerate(fila.cells):
        celda.width = anchos[j]; bordes_celda(celda); celda.paragraphs[0].text = ''
    p = fila.cells[0].paragraphs[0]; p.paragraph_format.space_after = Pt(0)
    txt(p, 'Otros:', size=9.5)
    der = fila.cells[1].merge(fila.cells[3])
    p = limpiar(der)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_after = Pt(0)
    txt(p, '{{NOTA_FISCAL}}', size=9.5)

    # ---------- 03 Condiciones ----------
    titulo_seccion(doc, '03', 'Condiciones')
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.space_after = Pt(6)
    txt(p, '{{CONDICION}}', size=10)

    parrafo(doc, '{{PARRAFO_CIERRE}}', size=10, antes=14, despues=0)

    # ---------- Pie de firma ----------
    doc.add_paragraph().paragraph_format.space_after = Pt(6)
    f = doc.add_table(rows=1, cols=3); f.alignment = WD_TABLE_ALIGNMENT.CENTER
    no_partir(f.rows[0])
    f.columns[0].width, f.columns[1].width, f.columns[2].width = Mm(64), Mm(4), Mm(102)
    c_marca, c_div, c_datos = f.rows[0].cells
    c_marca.width, c_div.width, c_datos.width = Mm(64), Mm(4), Mm(102)

    p = c_marca.paragraphs[0]; p.paragraph_format.space_after = Pt(8)
    p.add_run().add_picture(os.path.join(IMG, 'itsa_caja_azul_logo_blanco.png'), width=Mm(62))
    p = c_marca.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(os.path.join(IMG, 'logo_flexnet.png'), width=Mm(36))

    _borde(c_div._tc.get_or_add_tcPr(), 'left', AZUL_HEX, sz=8)

    c_datos.paragraphs[0].text = ''
    p = c_datos.paragraphs[0]; p.paragraph_format.space_after = Pt(1)
    txt(p, '{{FIRMANTE_NOMBRE}}', size=12, bold=True, color=AZUL)
    parrafo(c_datos, '{{FIRMANTE_CARGO}}', size=9, bold=True, color=RGBColor(0x33, 0x34, 0x34), despues=8)
    for campo in ('{{FIRMANTE_DIRECCION}}', '{{FIRMANTE_TELEFONO}}',
                  '{{FIRMANTE_CORREO}}', '{{FIRMANTE_WEBS}}'):
        parrafo(c_datos, campo, size=9, color=GRIS_TEXTO, despues=4)

    doc.save(salida)
    print('Plantilla Word generada:', salida)


if __name__ == '__main__':
    destino = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        RAIZ, 'cotizacion', 'PLANTILLA_COTIZACION_EN_BLANCO.docx')
    construir(destino)
