#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Motor de llenado de la plantilla de cotizaciones ITSA.

Toma una FICHA DE DATOS en JSON y la funde con la plantilla en blanco para
producir el documento final. La plantilla nunca se modifica: se copia y se
rellenan sus marcadores.

    python3 generar_cotizacion.py ficha.json                 -> HTML + PDF + DOCX
    python3 generar_cotizacion.py ficha.json --solo-html
    python3 generar_cotizacion.py ficha.json --salida /ruta/

Cómo se incrusta la información:
  1. Campos simples   {{CAMPO}}            -> reemplazo directo.
  2. Bloques repetidos INICIO_x / FIN_x    -> el bloque se clona una vez por
     registro, de modo que la cotización admite cualquier número de servicios
     (se agregan o se quitan filas sin tocar la maqueta).
  3. Totales          -> se calculan desde los ítems; si la ficha no los pide,
     la fila completa se elimina.
  4. Campos ausentes  -> se marcan como [PENDIENTE: campo] en vez de inventarse.
"""
import argparse, base64, json, locale, mimetypes, os, re, shutil, subprocess, sys
from datetime import date, datetime
from decimal import Decimal, ROUND_HALF_UP

RAIZ = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
PLANTILLA_HTML = os.path.join(RAIZ, 'cotizacion', 'plantilla_cotizacion.html')
PLANTILLA_DOCX = os.path.join(RAIZ, 'cotizacion', 'PLANTILLA_COTIZACION_EN_BLANCO.docx')

MESES = ['enero', 'febrero', 'marzo', 'abril', 'mayo', 'junio',
         'julio', 'agosto', 'septiembre', 'octubre', 'noviembre', 'diciembre']

CIERRE_POR_DEFECTO = ('Quedamos atentos a sus comentarios y a coordinar el inicio del '
                      'proceso. Trabajaremos en conjunto con ustedes para que la '
                      'operación fluya sin complicaciones.')


# ------------------------------------------------------------------ utilidades
def pendiente(campo):
    return '[PENDIENTE: %s]' % campo


def dato(ficha, ruta, por_defecto=None):
    """Lee 'cliente.nombre' dentro de la ficha; marca PENDIENTE si falta."""
    nodo = ficha
    for parte in ruta.split('.'):
        if not isinstance(nodo, dict) or parte not in nodo:
            return por_defecto if por_defecto is not None else pendiente(ruta)
        nodo = nodo[parte]
    if nodo in (None, '', []):
        return por_defecto if por_defecto is not None else pendiente(ruta)
    return nodo


def fecha_larga(iso):
    try:
        d = datetime.strptime(str(iso)[:10], '%Y-%m-%d').date()
    except (ValueError, TypeError):
        return str(iso)
    return '%d de %s de %d' % (d.day, MESES[d.month - 1], d.year)


def money(valor, simbolo='$'):
    if isinstance(valor, str):
        return valor
    q = Decimal(str(valor)).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    return '%s%s' % (simbolo, '{:,.2f}'.format(q))


def escapar(texto):
    return (str(texto).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;'))


# ------------------------------------------------------- normalización de ficha
def contexto(ficha):
    """Convierte la ficha de datos en el diccionario de campos de la plantilla."""
    doc = ficha.get('documento', {})
    hoy = doc.get('fecha_emision') or date.today().isoformat()
    moneda = doc.get('moneda', 'USD')
    simbolo = {'USD': '$', 'EUR': '€'}.get(moneda, '$')

    items = ficha.get('items') or []
    subtotal = Decimal('0')
    for it in items:
        try:
            subtotal += Decimal(str(it.get('total', 0)))
        except Exception:
            pass

    totales = ficha.get('totales') or {}
    mostrar_totales = bool(totales.get('mostrar', len(items) > 1))
    tasa_iva = Decimal(str(totales.get('tasa_iva', 0)))
    iva = (subtotal * tasa_iva / 100).quantize(Decimal('0.01'), rounding=ROUND_HALF_UP)
    total = subtotal + iva if tasa_iva else subtotal

    if tasa_iva:
        etiqueta_total, valor_total = 'Total (IVA %g %% incluido)' % float(tasa_iva), total
        nota = totales.get('nota', 'Valores incluyen IVA')
    else:
        etiqueta_total, valor_total = 'Total estimado', subtotal
        nota = totales.get('nota', 'Valores no incluyen IVA')

    cliente = dato(ficha, 'cliente.nombre')
    ciudades = dato(ficha, 'cliente.ciudades_operacion', [])
    if isinstance(ciudades, list):
        ciudades = ', '.join(ciudades) if ciudades else pendiente('cliente.ciudades_operacion')

    webs = dato(ficha, 'emisor.webs', [])
    if isinstance(webs, list):
        webs = '<br>'.join(webs) if webs else pendiente('emisor.webs')

    ctx = {
        'ETIQUETA_DOCUMENTO':  doc.get('etiqueta', 'Cotización'),
        'TITULO_DOCUMENTO':    doc.get('titulo', 'Cotización de servicios logísticos'),
        'CODIGO_DOCUMENTO':    dato(ficha, 'documento.codigo'),
        'VERSION':             doc.get('version', 'v1'),
        'REGION':              dato(ficha, 'emisor.region', 'Ecuador'),
        'CIUDAD_EMISION':      doc.get('ciudad', dato(ficha, 'emisor.region', 'Quito')),
        'FECHA_EMISION_LARGA': fecha_larga(hoy),
        'MONEDA':              moneda,
        'CLIENTE':             cliente,
        'CLIENTE_MAYUSCULAS':  str(cliente).upper(),
        'ATENCION_A':          dato(ficha, 'cliente.atencion_a'),
        'VALIDEZ':             doc.get('validez', '30 días calendario'),
        'GERENTE_COMERCIAL':   dato(ficha, 'emisor.responsable'),
        'SERVICIO_COTIZADO':   dato(ficha, 'servicio_cotizado'),
        'CIUDADES_OPERACION':  ciudades,
        'NOTA_FISCAL':         nota,
        'TOTAL_ETIQUETA':      etiqueta_total,
        'TOTAL_VALOR':         money(valor_total, simbolo),
        'PARRAFO_CIERRE':      ficha.get('parrafo_cierre') or CIERRE_POR_DEFECTO,
        'FIRMANTE_NOMBRE':     dato(ficha, 'emisor.responsable'),
        'FIRMANTE_CARGO':      dato(ficha, 'emisor.cargo', 'Gerente Comercial'),
        'FIRMANTE_DIRECCION':  dato(ficha, 'emisor.direccion'),
        'FIRMANTE_TELEFONO':   dato(ficha, 'emisor.telefono'),
        'FIRMANTE_CORREO':     dato(ficha, 'emisor.correo'),
        'FIRMANTE_WEBS':       webs,
    }

    filas = [{
        'ITEM_SERVICIO': it.get('servicio', pendiente('items[].servicio')),
        'ITEM_CANTIDAD': it.get('cantidad', ''),
        'ITEM_TARIFA':   it.get('tarifa', ''),
        'ITEM_TOTAL':    money(it.get('total', 0), simbolo),
    } for it in items] or [{'ITEM_SERVICIO': pendiente('items'), 'ITEM_CANTIDAD': '',
                            'ITEM_TARIFA': '', 'ITEM_TOTAL': ''}]

    condiciones = ficha.get('condiciones') or [pendiente('condiciones')]
    return ctx, filas, condiciones, mostrar_totales


# ----------------------------------------------------------------- render HTML
def bloque(html, nombre):
    patron = re.compile(r'[ \t]*<!--\s*INICIO_%s\s*-->(.*?)[ \t]*<!--\s*FIN_%s\s*-->\n?'
                        % (nombre, nombre), re.S)
    m = patron.search(html)
    if not m:
        raise SystemExit('La plantilla no contiene el bloque %s' % nombre)
    return patron, m.group(1)


def render_html(ficha, destino, standalone=False):
    with open(PLANTILLA_HTML, encoding='utf-8') as fh:
        html = fh.read()
    ctx, filas, condiciones, mostrar_totales = contexto(ficha)

    pat_items, molde = bloque(html, 'ITEMS')
    repetido = ''.join(
        re.sub(r'\{\{(\w+)\}\}', lambda m: escapar(f.get(m.group(1), '')), molde)
        for f in filas)
    html = pat_items.sub(lambda _: repetido, html)

    pat_tot, molde_tot = bloque(html, 'TOTALES')
    html = pat_tot.sub(lambda _: (molde_tot if mostrar_totales else ''), html)

    pat_cond, molde_cond = bloque(html, 'CONDICIONES')
    repetido = ''.join(molde_cond.replace('{{CONDICION}}', escapar(c)) for c in condiciones)
    html = pat_cond.sub(lambda _: repetido, html)

    html = re.sub(r'\{\{(\w+)\}\}',
                  lambda m: str(ctx.get(m.group(1), pendiente(m.group(1)))), html)

    if standalone:
        base = os.path.dirname(PLANTILLA_HTML)

        def incrustar(m):
            ruta = os.path.normpath(os.path.join(base, m.group(2)))
            if not os.path.exists(ruta):
                return m.group(0)
            tipo = mimetypes.guess_type(ruta)[0] or 'application/octet-stream'
            with open(ruta, 'rb') as fh:
                b64 = base64.b64encode(fh.read()).decode()
            return '%s(data:%s;base64,%s)' % (m.group(1), tipo, b64) \
                if m.group(1) == 'url' else '%s"data:%s;base64,%s"' % (m.group(1), tipo, b64)

        html = re.sub(r"(url)\('(\.\./[^']+)'\)", incrustar, html)
        html = re.sub(r'(src=)"(\.\./[^"]+)"', incrustar, html)

    with open(destino, 'w', encoding='utf-8') as fh:
        fh.write(html)
    return destino


def html_a_pdf(origen, destino):
    for exe in ('/opt/pw-browsers/chromium-1194/chrome-linux/chrome',
                shutil.which('chromium'), shutil.which('google-chrome'),
                shutil.which('chromium-browser')):
        if exe and os.path.exists(exe):
            subprocess.run([exe, '--headless', '--disable-gpu', '--no-sandbox',
                            '--allow-file-access-from-files', '--virtual-time-budget=8000',
                            '--no-pdf-header-footer', '--print-to-pdf=' + destino,
                            'file://' + os.path.abspath(origen)],
                           check=True, capture_output=True)
            return destino
    print('  (sin navegador disponible: abra el HTML e imprima a PDF)', file=sys.stderr)
    return None


# ----------------------------------------------------------------- render DOCX
def _sustituir_parrafo(parrafo, ctx):
    completo = ''.join(r.text for r in parrafo.runs)
    if '{{' not in completo:
        return
    nuevo = re.sub(r'\{\{(\w+)\}\}',
                   lambda m: str(ctx.get(m.group(1), pendiente(m.group(1)))),
                   completo).replace('<br>', ' / ')
    if nuevo == completo:
        return
    for extra in parrafo.runs[1:]:
        extra._element.getparent().remove(extra._element)
    if parrafo.runs:
        parrafo.runs[0].text = nuevo


def _recorrer(contenedor, ctx):
    for p in contenedor.paragraphs:
        _sustituir_parrafo(p, ctx)
    for t in contenedor.tables:
        for fila in t.rows:
            for celda in fila.cells:
                _recorrer(celda, ctx)


def render_docx(ficha, destino):
    import copy
    from docx import Document
    ctx, filas, condiciones, mostrar_totales = contexto(ficha)
    doc = Document(PLANTILLA_DOCX)

    tabla = doc.tables[1]                      # 0 = datos, 1 = detalle económico
    fila_molde = tabla.rows[1]                 # fila patrón del ítem
    fila_totales = tabla.rows[2]

    # 1 fila por servicio: así la cotización crece o se reduce sin tocar la maqueta
    anterior = fila_molde._tr
    for f in filas[1:]:
        clon = copy.deepcopy(fila_molde._tr)
        anterior.addnext(clon)
        anterior = clon
    for fila, valores in zip(tabla.rows[1:1 + len(filas)], filas):
        for celda in fila.cells:
            _recorrer(celda, valores)

    if not mostrar_totales:
        fila_totales._tr.getparent().remove(fila_totales._tr)

    # 1 viñeta por condición
    molde_cond = next(p for p in doc.paragraphs if '{{CONDICION}}' in p.text)
    anterior = molde_cond._p
    for c in condiciones[1:]:
        clon = copy.deepcopy(molde_cond._p)
        anterior.addnext(clon)
        anterior = clon
    for p, c in zip([p for p in doc.paragraphs if '{{CONDICION}}' in p.text], condiciones):
        _sustituir_parrafo(p, {'CONDICION': c})

    _recorrer(doc, ctx)
    for seccion in doc.sections:
        _recorrer(seccion.header, ctx)
        _recorrer(seccion.footer, ctx)

    doc.save(destino)
    return destino


# ------------------------------------------------------------------------ main
def nombre_archivo(ficha):
    doc = ficha.get('documento', {})
    cliente = re.sub(r'[^0-9A-Za-zÁÉÍÓÚÑáéíóúñ]+', '',
                     str(ficha.get('cliente', {}).get('nombre', 'Cliente')))
    return '%s_%s_%s' % (doc.get('codigo', 'COT-AAAA-NNN'),
                         cliente or 'Cliente', doc.get('version', 'v1'))


def main():
    ap = argparse.ArgumentParser(description='Genera la cotización ITSA desde una ficha JSON.')
    ap.add_argument('ficha')
    ap.add_argument('--salida', default='.')
    ap.add_argument('--solo-html', action='store_true')
    ap.add_argument('--vinculado', action='store_true',
                    help='no incrusta los recursos; el HTML apunta a ../assets '
                         '(solo sirve si la salida queda junto a la plantilla)')
    args = ap.parse_args()

    with open(args.ficha, encoding='utf-8') as fh:
        ficha = json.load(fh)

    os.makedirs(args.salida, exist_ok=True)
    base = os.path.join(args.salida, nombre_archivo(ficha))

    html = render_html(ficha, base + '.html', standalone=not args.vinculado)
    print('HTML :', html)
    if not args.solo_html:
        pdf = html_a_pdf(html, base + '.pdf')
        if pdf:
            print('PDF  :', pdf)
        try:
            print('DOCX :', render_docx(ficha, base + '.docx'))
        except Exception as e:                                   # pragma: no cover
            print('DOCX : no generado (%s)' % e, file=sys.stderr)

    faltantes = sorted(set(re.findall(r'\[PENDIENTE: ([^\]]+)\]', open(html, encoding='utf-8').read())))
    if faltantes:
        print('\nDatos pendientes de completar:')
        for f in faltantes:
            print('  -', f)


if __name__ == '__main__':
    main()
